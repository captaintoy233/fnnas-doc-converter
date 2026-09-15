"""PPTX → Markdown。python-pptx 提取幻灯片文本。

架构改进（v2 - Document Model）:
- 先解析为统一 Document Model，再通过共享序列化器输出 Markdown
- 支持表格、组合图形（递归）、演讲者备注

向后兼容:
- convert() 方法保持原有签名和返回类型
- 新增 convert_to_document() 返回 Document Model
"""
from pptx import Presentation
from . import BaseConverter

# Document Model 导入
from model.document import Document
from model.block import Heading, Paragraph as MdParagraph, Table as MdTable, BlockQuote
from model.inline import Text, Bold
from model.table import TableRow, TableCell
from render.markdown import document_to_markdown


def _shape_to_blocks(shape, blocks: list):
    """递归提取形状内容为 Document Model blocks"""
    if shape.shape_type == 6:  # GROUP
        for sub in shape.shapes:
            _shape_to_blocks(sub, blocks)
        return
    
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            t = para.text.strip()
            if t:
                blocks.append(MdParagraph(children=[Text(t)]))
    
    if getattr(shape, "has_table", False):
        try:
            tbl = shape.table
            rows = []
            for ri, row in enumerate(tbl.rows):
                cells = []
                for cell in row.cells:
                    c = (cell.text or "").strip()
                    cells.append(TableCell(
                        children=[Text(c)] if c else [],
                        is_header=(ri == 0)
                    ))
                if any(c.text.strip() for c in cells):
                    rows.append(TableRow(cells=cells, is_header=(ri == 0)))
            
            if rows:
                blocks.append(MdTable(rows=rows))
        except Exception:
            pass


class PPTXConverter(BaseConverter):
    """PPTX → MD (python-pptx提取每页文字/表格/备注)"""

    def supported_extensions(self) -> list:
        return ['.pptx']

    def convert_to_document(self, file_path: str, **kwargs) -> Document:
        """转换为统一文档模型"""
        prs = Presentation(file_path)
        result = Document()
        
        # 摘要标题
        result.add_block(Heading(level=1, children=[Text("幻灯片摘要")]))
        result.add_block(MdParagraph(children=[Text(f"共 {len(prs.slides)} 页")]))
        
        for i, slide in enumerate(prs.slides):
            # 每页标题
            result.add_block(Heading(level=2, children=[Text(f"第 {i+1} 页")]))
            
            # 提取形状内容
            slide_blocks = []
            for shape in slide.shapes:
                _shape_to_blocks(shape, slide_blocks)
            
            for block in slide_blocks:
                result.add_block(block)
            
            # 演讲者备注
            try:
                if slide.has_notes_slide:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        result.add_block(BlockQuote(blocks=[
                            MdParagraph(children=[Bold([Text("备注")]), Text(": " + notes)])
                        ]))
            except Exception:
                pass
        
        return result

    def convert(self, file_path: str, **kwargs) -> str:
        """转换为 Markdown 字符串（向后兼容接口）"""
        doc = self.convert_to_document(file_path, **kwargs)
        return document_to_markdown(doc)
